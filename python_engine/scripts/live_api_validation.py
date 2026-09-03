from __future__ import annotations

import argparse
import json
import time
from datetime import UTC, datetime
from pathlib import Path
from typing import Any

from fastapi.testclient import TestClient
from openpyxl import load_workbook

from landfinder.api.app import create_app
from landfinder.catalog.chile import LAND_SEARCH_KEYWORDS
from landfinder.core.settings import load_settings

try:
    from scripts.live_data_e2e import _apply_secrets, _load_runtime_secrets
except ModuleNotFoundError:  # pragma: no cover - depends on invocation cwd.
    from live_data_e2e import _apply_secrets, _load_runtime_secrets

REPO_ROOT = Path(__file__).resolve().parents[2]
TERMINAL_STATUSES = {
    "COMPLETED",
    "COMPLETED_WITH_WARNINGS",
    "COMPLETED_WITH_ERRORS",
    "FAILED",
    "CANCELLED",
}


def main() -> None:
    parser = argparse.ArgumentParser(description="Validate a live E2E through FastAPI.")
    parser.add_argument("--data-dir", type=Path)
    parser.add_argument(
        "--report",
        type=Path,
        default=REPO_ROOT / "docs" / "LIVE_E2E_REPORT.md",
    )
    args = parser.parse_args()

    settings = load_settings(data_dir=args.data_dir)
    _apply_secrets(settings, _load_runtime_secrets())
    app = create_app(settings)

    with TestClient(app) as client:
        token = client.get("/session").json()["token"]
        headers = {"Authorization": f"Bearer {token}"}
        diagnostics = _credential_statuses(client, headers)
        searches: list[dict[str, Any]] = []

        if diagnostics.get("mercado_publico") == "CONNECTED":
            searches.append(
                _run_search(
                    client,
                    headers,
                    {
                        "search_sources": ["mercado_publico"],
                        "enrichment_sources": ["cmf_uf", "minvu"],
                        "keywords": list(LAND_SEARCH_KEYWORDS),
                        "keyword_mode": "ANY",
                        "extra_filters": {"mercado_publico_max_detail_records": 50},
                    },
                )
            )

        if diagnostics.get("mercadolibre_real_estate") == "CONNECTED":
            searches.append(
                _run_search(
                    client,
                    headers,
                    {
                        "search_sources": ["mercadolibre_real_estate"],
                        "enrichment_sources": ["cmf_uf", "minvu"],
                        "keywords": ["terreno", "lote", "sitio eriazo", "peladero"],
                        "keyword_mode": "ANY",
                        "extra_filters": {"limit": 50},
                    },
                )
            )

        properties = client.get("/properties", headers=headers, params={"limit": 1000}).json()[
            "items"
        ]
        selected = _select_record_for_review(searches, properties)
        property_payload = None
        territory = None
        if selected and selected.get("property_id"):
            property_payload = client.get(
                f"/properties/{selected['property_id']}",
                headers=headers,
            ).json()
            territory = client.get(
                f"/properties/{selected['property_id']}/territory",
                headers=headers,
            ).json()

        excel = _export_excel(client, headers)
        summary = {
            "checked_at": datetime.now(UTC).isoformat(),
            "credentials": diagnostics,
            "searches": searches,
            "selected_source_record": selected,
            "selected_property": _property_summary(property_payload),
            "territory": _territory_summary(territory),
            "excel": excel,
            "completion": _completion(selected, property_payload, territory, excel),
        }
        _write_report(args.report, summary)
        print(json.dumps(_console_summary(summary), ensure_ascii=True, indent=2, default=str))


def _credential_statuses(client: TestClient, headers: dict[str, str]) -> dict[str, str]:
    response = client.post("/credentials/diagnostics", headers=headers)
    response.raise_for_status()
    statuses: dict[str, str] = {}
    for item in response.json()["items"]:
        statuses[item["source_id"]] = item["status"]
    return statuses


def _run_search(
    client: TestClient,
    headers: dict[str, str],
    payload: dict[str, Any],
) -> dict[str, Any]:
    created = client.post("/searches", headers=headers, json=payload)
    created.raise_for_status()
    search_id = created.json()["search_id"]
    job = _wait_for_job(client, headers, search_id)
    diagnostics = client.get(
        f"/searches/{search_id}/classifier-diagnostics",
        headers=headers,
    ).json()
    source_records = {
        "all": _source_records(client, headers, search_id, bucket=None),
        "accepted": _source_records(client, headers, search_id, bucket="accepted"),
        "possible": _source_records(client, headers, search_id, bucket="possible"),
        "rejected": _source_records(client, headers, search_id, bucket="rejected"),
    }
    return {
        "search_id": search_id,
        "requested": payload,
        "status": job.get("status"),
        "source_run_status": _compact_source_run_status(job.get("source_run_status", [])),
        "found_count": job.get("found_count"),
        "processed_count": job.get("processed_count"),
        "discarded_count": job.get("discarded_count"),
        "source_records_received": job.get("source_records_received"),
        "source_records_classified": job.get("source_records_classified"),
        "accepted_candidates": job.get("accepted_candidates"),
        "normalized_candidates": job.get("normalized_candidates"),
        "persisted_properties": job.get("persisted_properties"),
        "possible_candidates": job.get("possible_candidates"),
        "rejected_records": job.get("rejected_records"),
        "classifier_distribution": diagnostics.get("classifier_distribution"),
        "source_records": source_records,
        "events": _compact_events(job.get("events", [])),
    }


def _wait_for_job(
    client: TestClient,
    headers: dict[str, str],
    search_id: str,
    *,
    timeout_seconds: int = 180,
) -> dict[str, Any]:
    deadline = time.monotonic() + timeout_seconds
    while time.monotonic() < deadline:
        response = client.get(f"/searches/{search_id}", headers=headers)
        response.raise_for_status()
        payload = response.json()
        if payload.get("status") in TERMINAL_STATUSES:
            return payload
        time.sleep(0.75)
    raise TimeoutError(f"SearchJob {search_id} did not finish in {timeout_seconds}s")


def _source_records(
    client: TestClient,
    headers: dict[str, str],
    search_id: str,
    *,
    bucket: str | None,
) -> dict[str, Any]:
    params = {"limit": 25}
    if bucket:
        params["bucket"] = bucket
    response = client.get(
        f"/searches/{search_id}/source-records",
        headers=headers,
        params=params,
    )
    response.raise_for_status()
    payload = response.json()
    return {
        "total": payload.get("total"),
        "sample": [_source_record_summary(item) for item in payload.get("items", [])[:10]],
    }


def _source_record_summary(item: dict[str, Any]) -> dict[str, Any]:
    return {
        "source_record_id": item.get("source_record_id"),
        "job_id": item.get("job_id"),
        "source": item.get("source"),
        "external_id": item.get("external_id"),
        "title": item.get("title"),
        "organization": item.get("organization"),
        "organization_region": item.get("organization_region"),
        "organization_commune": item.get("organization_commune"),
        "organization_address": item.get("organization_address"),
        "property_region": item.get("property_region") or item.get("region"),
        "property_commune": item.get("property_commune") or item.get("commune"),
        "property_address": item.get("property_address"),
        "amount_raw": item.get("amount_raw"),
        "amount": item.get("amount"),
        "amount_type": item.get("amount_type"),
        "surface_raw": item.get("surface_raw"),
        "surface_m2": item.get("surface_m2"),
        "detail_fetched": item.get("detail_fetched"),
        "detail_attempts": item.get("detail_attempts"),
        "detail_http_status": item.get("detail_http_status"),
        "fields_extracted": item.get("fields_extracted"),
        "fields_missing": item.get("fields_missing"),
        "extraction_reasons": item.get("extraction_reasons"),
        "data_completeness_score": item.get("data_completeness_score"),
        "classification": item.get("automatic_classification"),
        "confidence": item.get("automatic_confidence"),
        "score": item.get("automatic_score"),
        "positive_terms": item.get("positive_terms"),
        "negative_terms": item.get("negative_terms"),
        "rules_triggered": item.get("rules_triggered"),
        "manual_classification": item.get("manual_classification"),
        "property_id": item.get("property_id"),
        "source_url": item.get("source_url"),
    }


def _select_record_for_review(
    searches: list[dict[str, Any]],
    properties: list[dict[str, Any]],
) -> dict[str, Any] | None:
    properties_by_id = {item.get("internal_id"): item for item in properties}
    candidates: list[dict[str, Any]] = []
    for search in searches:
        accepted = search.get("source_records", {}).get("accepted", {}).get("sample", [])
        possible = search.get("source_records", {}).get("possible", {}).get("sample", [])
        candidates.extend(accepted)
        candidates.extend(possible)
    if not candidates:
        return None

    def score(item: dict[str, Any]) -> tuple[int, float, int]:
        property_item = properties_by_id.get(item.get("property_id"))
        value = 100 if item.get("property_id") else 0
        if item.get("classification") in {"LAND_ACQUISITION", "LAND_SALE"}:
            value += 50
        if _has_value(property_item, "asking_price_clp") or _has_value(
            property_item,
            "asking_price_uf",
        ):
            value += 40
        if _has_value(property_item, "gross_area_m2"):
            value += 30
        if _has_value(property_item, "latitude") and _has_value(property_item, "longitude"):
            value += 30
        return value, float(item.get("confidence") or 0), int(item.get("score") or 0)

    return sorted(candidates, key=score, reverse=True)[0]


def _export_excel(client: TestClient, headers: dict[str, str]) -> dict[str, Any]:
    response = client.post("/exports/excel", headers=headers)
    response.raise_for_status()
    payload = response.json()
    workbook = load_workbook(payload["path"])
    sheets = workbook.sheetnames
    workbook.close()
    return {
        "path": payload["path"],
        "status": payload["status"],
        "openpyxl_valid": True,
        "sheets": sheets,
    }


def _compact_source_run_status(rows: list[dict[str, Any]]) -> list[dict[str, Any]]:
    return [
        {
            "source_id": row.get("source_id"),
            "source_type": row.get("source_type"),
            "enabled": row.get("enabled"),
            "configured": row.get("configured"),
            "attempted": row.get("attempted"),
            "status": row.get("status"),
            "records_received": row.get("records_received"),
            "records_accepted": row.get("records_accepted"),
            "records_rejected": row.get("records_rejected"),
        }
        for row in rows
    ]


def _compact_events(events: list[dict[str, Any]]) -> list[dict[str, Any]]:
    return [
        {
            "type": event.get("type"),
            "stage": event.get("stage"),
            "source": event.get("source"),
            "property_id": event.get("property_id"),
            "message": event.get("message"),
        }
        for event in events[-30:]
    ]


def _property_summary(item: dict[str, Any] | None) -> dict[str, Any] | None:
    if not item:
        return None
    return {
        "internal_id": item.get("internal_id"),
        "source": item.get("source"),
        "external_id": item.get("external_id"),
        "url": item.get("canonical_url"),
        "title": item.get("publication_title"),
        "price_clp": item.get("asking_price_clp"),
        "price_uf": item.get("asking_price_uf"),
        "uf_rate": item.get("uf_rate"),
        "uf_rate_date": item.get("uf_rate_date"),
        "uf_source": item.get("uf_source"),
        "uf_retrieved_at": item.get("uf_retrieved_at"),
        "uf_cache_hit": item.get("uf_cache_hit"),
        "area_m2": item.get("gross_area_m2"),
        "uf_m2": item.get("price_uf_m2_gross"),
        "region": item.get("region"),
        "commune": item.get("commune"),
        "address": item.get("normalized_address") or item.get("original_address"),
        "latitude": item.get("latitude"),
        "longitude": item.get("longitude"),
        "location_precision": item.get("location_precision"),
        "geometry_status": item.get("geometry_status"),
        "geometry_type": item.get("geometry_type"),
    }


def _territory_summary(item: dict[str, Any] | None) -> dict[str, Any] | None:
    if not item:
        return None
    return {
        "urban_limit": item.get("urban_limit"),
        "risk_intersection": item.get("risk_intersection"),
        "prc": item.get("prc"),
        "geometry_metrics": item.get("geometry_metrics"),
        "evidence_count": len(item.get("evidence") or []),
    }


def _completion(
    selected: dict[str, Any] | None,
    property_payload: dict[str, Any] | None,
    territory: dict[str, Any] | None,
    excel: dict[str, Any],
) -> dict[str, bool]:
    has_minvu_result = bool(
        territory
        and any(
            territory.get(key)
            for key in ("urban_limit", "risk_intersection", "prc", "geometry_metrics")
        )
    )
    return {
        "has_real_external_record": selected is not None,
        "has_persisted_property": property_payload is not None,
        "has_price": _has_value(property_payload, "asking_price_clp")
        or _has_value(property_payload, "asking_price_uf"),
        "has_area": _has_value(property_payload, "gross_area_m2"),
        "has_coordinates": _has_usable_coordinates(property_payload),
        "has_minvu": has_minvu_result,
        "has_evidence": bool(has_minvu_result and territory and territory.get("evidence")),
        "has_excel": bool(excel.get("openpyxl_valid")),
    }


def _has_usable_coordinates(item: dict[str, Any] | None) -> bool:
    if not _has_value(item, "latitude") or not _has_value(item, "longitude"):
        return False
    precision = str((item or {}).get("location_precision") or "UNKNOWN").upper()
    return precision not in {
        "UNKNOWN",
        "ORGANIZATION_LOCATION_ONLY",
        "NOT_VERIFIED",
    }


def _has_value(item: dict[str, Any] | None, key: str) -> bool:
    return item is not None and item.get(key) not in (None, "", 0)


def _console_summary(summary: dict[str, Any]) -> dict[str, Any]:
    return {
        "checked_at": summary["checked_at"],
        "credentials": summary["credentials"],
        "completion": summary["completion"],
        "selected_source_record": summary["selected_source_record"],
        "selected_property": summary["selected_property"],
        "territory": summary["territory"],
        "excel": summary["excel"],
    }


def _write_report(path: Path, summary: dict[str, Any]) -> None:
    path.parent.mkdir(parents=True, exist_ok=True)
    lines = [
        "# Live E2E Report",
        "",
        f"Fecha UTC: {summary['checked_at']}",
        "",
        "## API Local",
        "",
        "Esta validacion fue ejecutada contra la app FastAPI local con TestClient.",
        "No se incluyen secretos, tokens, tickets ni claves API.",
        "",
        "## Credenciales",
        "",
        "| Fuente | Estado |",
        "|---|---|",
        f"| Mercado Publico | `{summary['credentials'].get('mercado_publico')}` |",
        f"| CMF | `{summary['credentials'].get('cmf')}` |",
        f"| MercadoLibre | `{summary['credentials'].get('mercadolibre_real_estate')}` |",
        "",
        "## SearchJobs Reales Por API",
        "",
    ]
    if summary["searches"]:
        lines.extend(
            [
                "| SearchJob | Fuente | Estado | Fuente recibidos | "
                "Clasificados | Aceptados | Posibles | Rechazados | Persistidos |",
                "|---|---|---|---:|---:|---:|---:|---:|---:|",
            ]
        )
        for search in summary["searches"]:
            lines.append(
                f"| `{search['search_id']}` | {', '.join(search['requested']['search_sources'])} | "
                f"`{search['status']}` | {search.get('source_records_received')} | "
                f"{search.get('source_records_classified')} | "
                f"{search.get('accepted_candidates')} | "
                f"{search.get('possible_candidates')} | {search.get('rejected_records')} | "
                f"{search.get('persisted_properties')} |"
            )
    else:
        lines.append("No se ejecuto SearchJob porque no hubo credenciales conectadas.")

    lines.extend(["", "## Clasificador Y Registros Fuente", "", "```json"])
    lines.append(json.dumps(summary["searches"], ensure_ascii=True, indent=2, default=str))
    lines.extend(["```", "", "## Registro Seleccionado", "", "```json"])
    lines.append(
        json.dumps(summary["selected_source_record"], ensure_ascii=True, indent=2, default=str)
    )
    lines.extend(["```", "", "## Propiedad Persistida", "", "```json"])
    lines.append(json.dumps(summary["selected_property"], ensure_ascii=True, indent=2, default=str))
    lines.extend(["```", "", "## Territorio / MINVU", "", "```json"])
    lines.append(json.dumps(summary["territory"], ensure_ascii=True, indent=2, default=str))
    lines.append("```")
    if not summary["completion"].get("has_minvu"):
        lines.extend(
            [
                "",
                "Nota: cualquier evidencia `UNKNOWN` de geometria es trazabilidad de ausencia "
                "de dato, no una consulta efectiva a capas MINVU.",
            ]
        )
    lines.extend(["", "## Excel", "", "```json"])
    lines.append(json.dumps(summary["excel"], ensure_ascii=True, indent=2, default=str))
    lines.extend(["```", "", "## Criterio De Termino", "", "```json"])
    lines.append(json.dumps(summary["completion"], ensure_ascii=True, indent=2, default=str))
    lines.append("```")
    path.write_text("\n".join(lines) + "\n", encoding="utf-8")


if __name__ == "__main__":
    main()
