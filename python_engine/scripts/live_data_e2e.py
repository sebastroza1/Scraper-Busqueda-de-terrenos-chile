from __future__ import annotations

import argparse
import asyncio
import ctypes
import json
import os
import re
import unicodedata
from ctypes import wintypes
from datetime import UTC, date, datetime
from pathlib import Path
from typing import Any

from openpyxl import load_workbook

from landfinder.api.jobs import JobManager, ensure_data_dirs, job_to_dict
from landfinder.catalog.chile import LAND_SEARCH_KEYWORDS
from landfinder.connectors.base import SearchQuery
from landfinder.connectors.public_procurement.mercado_publico import (
    classify_public_procurement_land_tender,
)
from landfinder.core.credentials import diagnose_source_credentials
from landfinder.core.settings import Settings, load_settings
from landfinder.core.uf import UFService, UFUnavailableError
from landfinder.geospatial.ipt_catalog import MinvuIptCatalog
from landfinder.geospatial.minvu import MinvuArcGisConnector
from landfinder.geospatial.territorial import TerritorialAnalyzer
from landfinder.reports.excel import export_properties_excel
from landfinder.storage.repositories import (
    Database,
    GeospatialCacheRepository,
    PropertyRepository,
)

SECRET_KEYS = (
    "MERCADO_PUBLICO_TICKET",
    "CMF_API_KEY",
    "MERCADOLIBRE_ACCESS_TOKEN",
    "MERCADOLIBRE_REFRESH_TOKEN",
    "MERCADOLIBRE_ACCESS_TOKEN_EXPIRES_AT",
)

REPO_ROOT = Path(__file__).resolve().parents[2]

SEARCH_TERMS = tuple(LAND_SEARCH_KEYWORDS)


async def main() -> None:
    parser = argparse.ArgumentParser(
        description="Run a live Data+GIS E2E without printing secrets."
    )
    parser.add_argument("--data-dir", type=Path)
    parser.add_argument(
        "--report",
        type=Path,
        default=REPO_ROOT / "docs" / "LIVE_E2E_REPORT.md",
    )
    args = parser.parse_args()

    settings = load_settings(data_dir=args.data_dir)
    secrets = _load_runtime_secrets()
    _apply_secrets(settings, secrets)
    ensure_data_dirs(settings.data_dir)

    database = Database(settings.database_path)
    diagnostics = await _diagnose(settings)
    uf = await _fetch_real_uf(settings) if diagnostics["cmf"] == "CONNECTED" else None
    run_started_at = datetime.now(UTC)
    jobs = await _run_search_jobs(settings, diagnostics)

    properties = PropertyRepository(database).list(limit=1000)
    reference_selection_note = "current_live_searchjob"
    reference = _select_reference_property(properties, seen_after=run_started_at)
    if reference is None and _jobs_rate_limited(jobs):
        reference = _select_reference_property(properties)
        if reference is not None:
            reference_selection_note = (
                "latest_persisted_live_record_reused_after_current_rate_limit"
            )
    territory: dict[str, Any] | None = None
    location_precision: dict[str, Any] | None = None
    if reference is not None:
        location_precision = _apply_location_precision(database, reference)
        reference = PropertyRepository(database).get(reference["internal_id"])
        if reference and _has_coordinates(reference):
            cache = GeospatialCacheRepository(database)
            minvu = MinvuArcGisConnector(cache=cache)
            analyzer = TerritorialAnalyzer(
                database=database,
                minvu=minvu,
                ipt_catalog=MinvuIptCatalog(minvu=minvu, cache=cache),
            )
            territory = await analyzer.refresh_property(
                reference["internal_id"],
                force_refresh=True,
            )
            reference = PropertyRepository(database).get(reference["internal_id"])

    excel = _export_and_validate(settings, database) if reference is not None else None
    summary = _summary(
        settings=settings,
        diagnostics=diagnostics,
        uf=uf,
        jobs=jobs,
        properties=PropertyRepository(database).list(limit=1000),
        reference=reference,
        location_precision=location_precision,
        territory=territory,
        excel=excel,
        reference_selection_note=reference_selection_note if reference else None,
    )
    _write_reference_file(settings.data_dir, reference)
    _write_report(args.report, summary)
    print(json.dumps(_console_summary(summary), ensure_ascii=True, indent=2, default=str))


async def _diagnose(settings: Settings) -> dict[str, str]:
    result: dict[str, str] = {}
    for source_id in ("mercado_publico", "cmf", "mercadolibre_real_estate"):
        diagnostic = await diagnose_source_credentials(settings, source_id)
        result[source_id if source_id != "mercadolibre_real_estate" else "mercadolibre"] = (
            diagnostic.status.value
        )
    return result


async def _fetch_real_uf(settings: Settings) -> dict[str, Any] | None:
    try:
        uf = await UFService(settings.database_path, settings.cmf_api_key).get_uf(date.today())
    except (UFUnavailableError, Exception) as exc:  # noqa: BLE001 - report sanitized state.
        return {"status": "UNAVAILABLE", "error": type(exc).__name__}
    return {
        "status": "CONNECTED",
        "date": uf.date.isoformat(),
        "value": str(uf.value),
        "source": uf.source,
        "retrieved_at": uf.retrieved_at.isoformat(),
    }


async def _run_search_jobs(settings: Settings, diagnostics: dict[str, str]) -> list[dict[str, Any]]:
    jobs = JobManager(settings)
    results = []
    if diagnostics.get("mercado_publico") == "CONNECTED":
        results.append(
            await _run_job(
                jobs,
                SearchQuery(
                    search_sources=["mercado_publico"],
                    enrichment_sources=["cmf_uf", "minvu"],
                    keywords=list(SEARCH_TERMS),
                    keyword_mode="ANY",
                    extra_filters={"mercado_publico_max_detail_records": 50},
                ),
            )
        )
    if diagnostics.get("mercadolibre") == "CONNECTED":
        results.append(
            await _run_job(
                jobs,
                SearchQuery(
                    search_sources=["mercadolibre_real_estate"],
                    enrichment_sources=["cmf_uf", "minvu"],
                    keywords=["terreno", "lote", "sitio eriazo", "peladero"],
                    keyword_mode="ANY",
                    extra_filters={"limit": 50},
                ),
            )
        )
    return results


async def _run_job(jobs: JobManager, query: SearchQuery) -> dict[str, Any]:
    job = jobs.create_search(query)
    if job.task is not None:
        await job.task
    payload = job_to_dict(job)
    return {
        "job_id": payload["id"],
        "status": str(payload["status"]),
        "sources": payload["sources"],
        "found_count": payload["found_count"],
        "processed_count": payload["processed_count"],
        "discarded_count": payload["discarded_count"],
        "warning_count": payload["warning_count"],
        "error_count": payload["error_count"],
        "source_run_status": payload.get("source_run_status", []),
        "events": [
            {
                "type": event.get("type"),
                "stage": event.get("stage"),
                "source": event.get("source"),
                "message": _sanitize_message(event.get("message")),
            }
            for event in payload["events"]
        ],
    }


def _select_reference_property(
    properties: list[dict[str, Any]],
    *,
    seen_after: datetime | None = None,
) -> dict[str, Any] | None:
    candidates = [
        item
        for item in properties
        if item.get("source") in {"mercadolibre_real_estate", "mercado_publico"}
        and item.get("external_id")
        and (seen_after is None or _seen_after(item, seen_after))
    ]
    if not candidates:
        return None

    def score(item: dict[str, Any]) -> tuple[int, str]:
        value = _semantic_reference_score(item)
        value += 50 if _is_meaningful_price(item) else 0
        value += 40 if item.get("gross_area_m2") else 0
        value += 35 if _has_coordinates(item) else 0
        value += 20 if item.get("canonical_url") else 0
        value += 10 if item.get("commune") else 0
        value += 10 if item.get("region") else 0
        value += 5 if item.get("source") == "mercadolibre_real_estate" else 0
        return value, str(item.get("last_seen") or "")

    return sorted(candidates, key=score, reverse=True)[0]


def _apply_location_precision(
    database: Database,
    reference: dict[str, Any],
) -> dict[str, Any]:
    if not _has_raw_coordinates(reference):
        precision = "UNKNOWN"
        reason = "No coordinates were returned by the external source."
        confidence = 0
    elif reference.get("source") == "mercadolibre_real_estate":
        precision = "APPROXIMATE_PROPERTY"
        reason = (
            "MercadoLibre returned item coordinates, but the current audit has not verified "
            "that they represent an exact parcel point."
        )
        confidence = 0.45
    else:
        precision = "UNKNOWN"
        reason = "The source did not provide evidence that coordinates are exact."
        confidence = 0.2

    with database.transaction() as connection:
        connection.execute(
            """
            UPDATE property_candidates
            SET location_precision = ?, geocoding_source = ?, geocoding_confidence = ?
            WHERE internal_id = ?
            """,
            (
                precision,
                str(reference.get("source") or "external_source"),
                confidence,
                reference["internal_id"],
            ),
        )
    reference["location_precision"] = precision
    reference["geocoding_source"] = str(reference.get("source") or "external_source")
    reference["geocoding_confidence"] = confidence
    return {"precision": precision, "reason": reason, "confidence": confidence}


def _export_and_validate(settings: Settings, database: Database) -> dict[str, Any]:
    output = Path(settings.data_dir) / "exports" / "landfinder_live_e2e.xlsx"
    export_properties_excel(database.database_path, output)
    workbook = load_workbook(output)
    sheets = workbook.sheetnames
    workbook.close()
    return {"path": str(output), "status": "created", "openpyxl_valid": True, "sheets": sheets}


def _summary(
    *,
    settings: Settings,
    diagnostics: dict[str, str],
    uf: dict[str, Any] | None,
    jobs: list[dict[str, Any]],
    properties: list[dict[str, Any]],
    reference: dict[str, Any] | None,
    location_precision: dict[str, Any] | None,
    territory: dict[str, Any] | None,
    excel: dict[str, Any] | None,
    reference_selection_note: str | None,
) -> dict[str, Any]:
    return {
        "checked_at": datetime.now(UTC).isoformat(),
        "data_dir": str(settings.data_dir),
        "database_path": str(settings.database_path),
        "credentials": diagnostics,
        "uf": uf or {"status": "NOT_CONFIGURED"},
        "jobs": jobs,
        "source_run_status": [
            row for job in jobs for row in job.get("source_run_status", [])
        ],
        "properties_total": len(properties),
        "reference_property": _reference_payload(reference),
        "reference_selection_note": reference_selection_note,
        "location_precision": location_precision,
        "territory": _territory_payload(territory),
        "excel": excel,
        "completion": {
            "has_real_external_record": reference is not None,
            "has_price": bool(
                reference and _is_meaningful_price(reference)
            ),
            "has_area": bool(reference and reference.get("gross_area_m2")),
            "has_real_uf": bool(uf and uf.get("status") == "CONNECTED"),
            "has_coordinates": bool(reference and _has_coordinates(reference)),
            "has_minvu": territory is not None,
            "has_evidence": bool(territory and territory.get("evidence")),
            "has_excel": bool(excel and excel.get("openpyxl_valid")),
        },
    }


def _reference_payload(reference: dict[str, Any] | None) -> dict[str, Any] | None:
    if reference is None:
        return None
    return {
        "internal_id": reference.get("internal_id"),
        "source": reference.get("source"),
        "external_id": reference.get("external_id"),
        "url": reference.get("canonical_url"),
        "title": reference.get("publication_title"),
        "price_clp": reference.get("asking_price_clp"),
        "price_uf": reference.get("asking_price_uf"),
        "area_m2": reference.get("gross_area_m2"),
        "uf_m2": reference.get("price_uf_m2_gross"),
        "region": reference.get("region"),
        "commune": reference.get("commune"),
        "address": reference.get("normalized_address") or reference.get("original_address"),
        "latitude": reference.get("latitude"),
        "longitude": reference.get("longitude"),
        "location_precision": reference.get("location_precision"),
        "procurement_classification": _reference_classification(reference),
    }


def _territory_payload(territory: dict[str, Any] | None) -> dict[str, Any] | None:
    if territory is None:
        return None
    urban_limit = territory.get("urban_limit") or {}
    risk = territory.get("risk_intersection") or {}
    prc = territory.get("prc") or {}
    metrics = territory.get("geometry_metrics") or {}
    return {
        "urban_limit": {
            "status": urban_limit.get("status"),
            "service": urban_limit.get("service"),
            "layer": urban_limit.get("layer"),
            "confidence": urban_limit.get("confidence"),
            "feature_count": urban_limit.get("feature_count"),
        },
        "risk_intersection": {
            "status": risk.get("status"),
            "risk_types": risk.get("risk_types"),
            "layers": risk.get("layers"),
            "confidence": risk.get("confidence"),
            "analysis_precision": risk.get("analysis_precision"),
        },
        "prc": {
            "status": prc.get("status"),
            "applicable_service": prc.get("applicable_service"),
            "applicable_layer": prc.get("applicable_layer"),
            "zone": prc.get("zone"),
            "zone_field": prc.get("zone_field"),
            "urban_zone": prc.get("urban_zone"),
        },
        "geometry_metrics": {
            "geometry_status": metrics.get("geometry_status"),
            "geometry_type": metrics.get("geometry_type"),
            "geometry_source": metrics.get("geometry_source"),
            "geometric_area_m2": metrics.get("geometric_area_m2"),
            "perimeter_m": metrics.get("perimeter_m"),
        },
        "evidence_count": len(territory.get("evidence") or []),
    }


def _console_summary(summary: dict[str, Any]) -> dict[str, Any]:
    return {
        "checked_at": summary["checked_at"],
        "credentials": summary["credentials"],
        "completion": summary["completion"],
        "reference_property": summary["reference_property"],
        "reference_selection_note": summary["reference_selection_note"],
        "territory": summary["territory"],
        "excel": summary["excel"],
    }


def _write_report(path: Path, summary: dict[str, Any]) -> None:
    path.parent.mkdir(parents=True, exist_ok=True)
    reference = summary["reference_property"]
    territory = summary["territory"]
    jobs = summary["jobs"]
    source_run_status = summary.get("source_run_status", [])
    lines = [
        "# Live E2E Report",
        "",
        f"Fecha UTC: {summary['checked_at']}",
        "",
        "## Credenciales",
        "",
        "| Fuente | Estado |",
        "|---|---|",
        f"| Mercado Publico | `{summary['credentials'].get('mercado_publico')}` |",
        f"| CMF | `{summary['credentials'].get('cmf')}` |",
        f"| MercadoLibre | `{summary['credentials'].get('mercadolibre')}` |",
        "",
        "## Fuentes Por Job",
        "",
    ]
    if source_run_status:
        lines.extend(
            [
                "| Fuente | Tipo | Habilitada | Configurada | Estado | "
                "Recibidos | Aceptados | Rechazados |",
                "|---|---|---:|---:|---|---:|---:|---:|",
            ]
        )
        for row in source_run_status:
            lines.append(
                f"| {row.get('source_id')} | {row.get('source_type')} | "
                + f"{_yes_no(row.get('enabled'))} | {_yes_no(row.get('configured'))} | "
                + f"`{row.get('status')}` | {row.get('records_received')} | "
                + f"{row.get('records_accepted')} | {row.get('records_rejected')} |"
            )
    else:
        lines.append("No hubo estados de fuente porque no se ejecuto ningun SearchJob.")
    lines.extend(["", "## SearchJobs Reales", ""])
    if jobs:
        lines.extend(
            [
                "| Fuentes | Estado | Recibidos | Persistidos | Descartados | Warnings | Errores |",
                "|---|---|---:|---:|---:|---:|---:|",
            ]
        )
        for job in jobs:
            lines.append(
                "| "
                + ", ".join(job["sources"])
                + f" | `{job['status']}` | {job['found_count']} | {job['processed_count']} | "
                + f"{job['discarded_count']} | {job['warning_count']} | {job['error_count']} |"
            )
    else:
        lines.append("No se ejecuto ningun SearchJob porque no hubo credenciales conectadas.")
    lines.extend(["", "## UF Real", "", "```json"])
    lines.append(json.dumps(summary["uf"], ensure_ascii=True, indent=2, default=str))
    lines.extend(["```", "", "## ReferenceProperty", ""])
    if summary.get("reference_selection_note"):
        lines.extend([f"Seleccion: `{summary['reference_selection_note']}`", ""])
    if reference:
        lines.extend(
            [
                "```json",
                json.dumps(reference, ensure_ascii=True, indent=2, default=str),
                "```",
            ]
        )
    else:
        lines.append("No se pudo seleccionar un registro externo real persistido.")
    lines.extend(["", "## Precision Geografica", "", "```json"])
    lines.append(
        json.dumps(summary["location_precision"], ensure_ascii=True, indent=2, default=str)
    )
    lines.extend(["```", "", "## MINVU Sobre El Mismo Registro", ""])
    if territory:
        lines.extend(
            [
                "```json",
                json.dumps(territory, ensure_ascii=True, indent=2, default=str),
                "```",
            ]
        )
    else:
        lines.append("No se ejecuto MINVU porque no hubo ReferenceProperty con coordenadas reales.")
    lines.extend(["", "## Excel", "", "```json"])
    lines.append(json.dumps(summary["excel"], ensure_ascii=True, indent=2, default=str))
    lines.extend(["```", "", "## Criterio De Termino", "", "```json"])
    lines.append(json.dumps(summary["completion"], ensure_ascii=True, indent=2, default=str))
    lines.extend(
        [
            "```",
            "",
            "## Warnings Y Errores",
            "",
            "Mensajes sanitizados de SearchJob:",
            "",
            "```json",
            json.dumps([job.get("events", []) for job in jobs], ensure_ascii=True, indent=2),
            "```",
            "",
            "No se incluyen secretos, tokens, tickets ni claves API.",
        ]
    )
    path.write_text("\n".join(lines) + "\n", encoding="utf-8")


def _write_reference_file(data_dir: Path, reference: dict[str, Any] | None) -> None:
    if reference is None:
        return
    path = data_dir / "reference_property.json"
    path.write_text(
        json.dumps(_reference_payload(reference), ensure_ascii=True, indent=2, default=str),
        encoding="utf-8",
    )


def _load_runtime_secrets() -> dict[str, str]:
    result: dict[str, str] = {}
    dev_values = _read_dev_secret_file()
    for key in SECRET_KEYS:
        value = os.getenv(key) or _read_windows_credential(key) or dev_values.get(key)
        if value:
            result[key] = value
    return result


def _apply_secrets(settings: Settings, values: dict[str, str]) -> None:
    for key, value in values.items():
        os.environ[key] = value
    settings.cmf_api_key = values.get("CMF_API_KEY") or os.getenv("UF_CMF_API_KEY")
    settings.mercado_publico_ticket = values.get("MERCADO_PUBLICO_TICKET")
    settings.mercadolibre_access_token = values.get("MERCADOLIBRE_ACCESS_TOKEN")
    settings.mercadolibre_refresh_token = values.get("MERCADOLIBRE_REFRESH_TOKEN")
    settings.mercadolibre_access_token_expires_at = values.get(
        "MERCADOLIBRE_ACCESS_TOKEN_EXPIRES_AT"
    )


def _read_dev_secret_file() -> dict[str, str]:
    roots = [
        os.getenv("LOCALAPPDATA"),
        os.getenv("APPDATA"),
    ]
    for root in roots:
        if not root:
            continue
        path = Path(root) / "Landfinder" / "dev_secrets.json"
        if not path.exists():
            continue
        try:
            payload = json.loads(path.read_text(encoding="utf-8"))
        except (OSError, json.JSONDecodeError):
            continue
        if isinstance(payload, dict):
            return {
                str(key): str(value)
                for key, value in payload.items()
                if value not in (None, "")
            }
    return {}


def _read_windows_credential(key: str) -> str | None:
    if os.name != "nt":
        return None
    target = f"LandfinderChile:{key}"
    credential_ptr = ctypes.c_void_p()
    if not ctypes.windll.advapi32.CredReadW(target, 1, 0, ctypes.byref(credential_ptr)):
        return None
    try:
        credential = ctypes.cast(credential_ptr, ctypes.POINTER(_Credential)).contents
        if not credential.CredentialBlob or credential.CredentialBlobSize == 0:
            return None
        raw = ctypes.string_at(credential.CredentialBlob, credential.CredentialBlobSize)
        return raw.decode("utf-8")
    finally:
        ctypes.windll.advapi32.CredFree(credential_ptr)


class _Credential(ctypes.Structure):
    _fields_ = [
        ("Flags", wintypes.DWORD),
        ("Type", wintypes.DWORD),
        ("TargetName", wintypes.LPWSTR),
        ("Comment", wintypes.LPWSTR),
        ("LastWritten", wintypes.FILETIME),
        ("CredentialBlobSize", wintypes.DWORD),
        ("CredentialBlob", ctypes.c_void_p),
        ("Persist", wintypes.DWORD),
        ("AttributeCount", wintypes.DWORD),
        ("Attributes", ctypes.c_void_p),
        ("TargetAlias", wintypes.LPWSTR),
        ("UserName", wintypes.LPWSTR),
    ]


def _has_coordinates(reference: dict[str, Any]) -> bool:
    if not _has_raw_coordinates(reference):
        return False
    precision = str(reference.get("location_precision") or "UNKNOWN").upper()
    return precision not in {
        "UNKNOWN",
        "ORGANIZATION_LOCATION_ONLY",
        "NOT_VERIFIED",
    }


def _has_raw_coordinates(reference: dict[str, Any]) -> bool:
    return _as_float(reference.get("latitude")) is not None and _as_float(
        reference.get("longitude")
    ) is not None


def _as_float(value: Any) -> float | None:
    try:
        return float(value)
    except (TypeError, ValueError):
        return None


def _semantic_reference_score(item: dict[str, Any]) -> int:
    if item.get("source") == "mercadolibre_real_estate":
        return 140

    text = _normalized_reference_text(item)
    score = 0
    has_land = bool(
        re.search(r"\b(terreno|terrenos|predio|predios|lote|macrolote)\b", text)
    )
    has_property = bool(
        re.search(r"\b(inmueble|inmuebles|propiedad|bien raiz|bienes raices)\b", text)
    )
    has_sale = bool(re.search(r"\b(venta|vender|enajenacion|remate|subasta)\b", text))
    has_acquisition = bool(
        re.search(r"\b(adquisicion|adquisici|compra|comprar|adquirir)\b", text)
    )
    if has_sale and (has_land or has_property):
        score += 130
    elif has_acquisition and (has_land or has_property):
        score += 125
    elif has_land:
        score += 70
    elif has_property:
        score += 40

    false_positive_terms = [
        "visita a terreno",
        "trabajo en terreno",
        "todoterreno",
        "obra civil",
        "obras civiles",
        "construccion plaza",
        "mantenimiento",
        "seguro",
        "poliza",
        "taludes",
        "guardia de seguridad",
    ]
    if any(term in text for term in false_positive_terms) and not (has_sale or has_acquisition):
        score -= 100
    return score


def _normalized_reference_text(item: dict[str, Any]) -> str:
    value = " ".join(
        str(part or "")
        for part in (
            item.get("publication_title"),
            item.get("commune"),
            item.get("region"),
            item.get("normalized_address"),
            item.get("original_address"),
        )
    )
    decomposed = unicodedata.normalize("NFD", value.lower())
    without_accents = "".join(
        char
        for char in decomposed
        if unicodedata.category(char) != "Mn" and char.isascii()
    )
    return " ".join(without_accents.split())


def _is_meaningful_price(reference: dict[str, Any]) -> bool:
    for key in ("asking_price_uf", "asking_price_clp"):
        value = _as_float(reference.get(key))
        if value is not None and value > 1000:
            return True
    return False


def _seen_after(item: dict[str, Any], seen_after: datetime) -> bool:
    last_seen = item.get("last_seen")
    if not last_seen:
        return False
    try:
        parsed = datetime.fromisoformat(str(last_seen))
    except ValueError:
        return False
    if parsed.tzinfo is None:
        parsed = parsed.replace(tzinfo=UTC)
    return parsed >= seen_after


def _reference_classification(reference: dict[str, Any]) -> str | None:
    for source_record in reference.get("source_records") or []:
        raw = source_record.get("raw_payload_json")
        if not raw:
            continue
        try:
            payload = json.loads(raw)
        except json.JSONDecodeError:
            continue
        if reference.get("source") == "mercado_publico":
            return classify_public_procurement_land_tender(payload).value
        landfinder = payload.get("landfinder") or {}
        classification = landfinder.get("procurement_classification")
        if classification:
            return str(classification)
    return None


def _jobs_rate_limited(jobs: list[dict[str, Any]]) -> bool:
    return any(
        row.get("status") == "RATE_LIMITED"
        for job in jobs
        for row in job.get("source_run_status", [])
    )


def _sanitize_message(value: Any) -> str:
    text = str(value or "")
    for key in SECRET_KEYS:
        secret = os.getenv(key)
        if secret:
            text = text.replace(secret, "[REDACTED]")
    return text


def _yes_no(value: Any) -> str:
    if value in (1, True, "1", "true", "True"):
        return "Si"
    if value in (0, False, "0", "false", "False"):
        return "No"
    return "-"


if __name__ == "__main__":
    asyncio.run(main())
