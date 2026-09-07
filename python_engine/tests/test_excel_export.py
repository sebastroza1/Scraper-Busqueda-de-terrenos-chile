from decimal import Decimal

from openpyxl import load_workbook

from landfinder.models.enums import SourceRole, SourceRunStatus, SourceType
from landfinder.models.property import DataProvenance
from landfinder.models.records import NormalizedPropertyRecord
from landfinder.reports.excel import export_properties_excel
from landfinder.storage.repositories import (
    Database,
    DataProvenanceRepository,
    PropertyRepository,
    SourceRecordAuditRepository,
    SourceRunStatusRepository,
)


def test_excel_export_opens(tmp_path) -> None:
    db = tmp_path / "landfinder.sqlite"
    database = Database(db)
    PropertyRepository(database).upsert(
        NormalizedPropertyRecord(
            source_id="fixture",
            source_type=SourceType.MARKETPLACE,
            external_id="abc",
            title="Terreno",
            gross_area_m2=Decimal("10000"),
            price_uf=Decimal("1000"),
            price_uf_m2=Decimal("0.10"),
        )
    )
    SourceRunStatusRepository(database).upsert(
        job_id="job-1",
        source_id="mercado_publico",
        source_type=SourceRole.SEARCH.value,
        enabled=True,
        configured=True,
        attempted=True,
        status=SourceRunStatus.SUCCESS.value,
        records_received=1,
        records_accepted=1,
        queries=["terreno"],
    )
    SourceRunStatusRepository(database).upsert(
        job_id="job-1",
        source_id="mercadolibre_real_estate",
        source_type=SourceRole.SEARCH.value,
        enabled=True,
        configured=False,
        attempted=False,
        status=SourceRunStatus.SKIPPED_NOT_CONFIGURED.value,
    )
    DataProvenanceRepository(database).upsert(
        DataProvenance(
            property_id=PropertyRepository(database).list()[0]["internal_id"],
            field_name="uf_rate",
            value="40000",
            source_id="cmf_uf",
            source_type=SourceRole.ENRICHMENT.value,
            method="official_api",
            source_url="https://api.cmfchile.cl/uf?apikey=secret-value",
            confidence=0.95,
        )
    )
    SourceRecordAuditRepository(database).upsert(
        "job-1",
        {
            "source": "mercado_publico",
            "source_type": SourceType.PUBLIC_PROCUREMENT.value,
            "external_id": "1234-1-LQ26",
            "title": "Compra terreno",
            "organization": "SERVIU",
            "source_url": "https://api.mercadopublico.cl/licitaciones?ticket=secret-value",
            "automatic_classification": "LAND_ACQUISITION",
            "automatic_confidence": 0.91,
            "automatic_score": 110,
            "classification_reason": "rules=TITLE_LAND_ACQUISITION",
            "positive_terms": ["compra", "terreno"],
            "negative_terms": [],
            "rules_triggered": ["TITLE_LAND_ACQUISITION"],
            "raw_payload": {"ticket": "secret-value", "Nombre": "Compra terreno"},
        },
    )

    output = export_properties_excel(db, tmp_path / "out.xlsx")
    workbook = load_workbook(output)
    assert {
        "Resumen",
        "Terrenos",
        "Fuentes_Utilizadas",
        "ANALISIS_TERRITORIAL",
        "CAPAS_MINVU",
        "Fuentes",
        "Historial_Precios",
        "Metricas_GIS",
        "Fuentes_Evidencia",
        "Classifier_Audit",
        "SourceRecords",
        "Errores",
    }.issubset(set(workbook.sheetnames))
    evidence_sheet = workbook["Fuentes_Evidencia"]
    evidence_values = [
        cell.value
        for row in evidence_sheet.iter_rows()
        for cell in row
        if isinstance(cell.value, str)
    ]
    assert any("apikey=REDACTED" in value for value in evidence_values)
    assert all("secret-value" not in value for value in evidence_values)
    classifier_values = [
        cell.value
        for row in workbook["Classifier_Audit"].iter_rows()
        for cell in row
        if isinstance(cell.value, str)
    ]
    assert any("ticket=REDACTED" in value for value in classifier_values)
    assert all("secret-value" not in value for value in classifier_values)
    assert workbook["SourceRecords"].max_row >= 2
    workbook.close()
