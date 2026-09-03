from __future__ import annotations

import json
import re
from datetime import UTC, datetime
from pathlib import Path
from typing import Any

from openpyxl import Workbook, load_workbook
from openpyxl.drawing.image import Image as ExcelImage
from openpyxl.styles import Font, PatternFill
from openpyxl.utils import get_column_letter

from landfinder import __version__
from landfinder.storage.database import connect


def export_properties_excel(database_path: str | Path, output_path: str | Path) -> Path:
    output = Path(output_path)
    output.parent.mkdir(parents=True, exist_ok=True)

    with connect(database_path) as connection:
        properties = [
            dict(row)
            for row in connection.execute(
                "SELECT * FROM property_candidates ORDER BY last_seen DESC"
            ).fetchall()
        ]
        sources = [
            dict(row)
            for row in connection.execute(
                "SELECT * FROM property_source_records ORDER BY last_seen DESC"
            ).fetchall()
        ]
        price_history = [
            dict(row)
            for row in connection.execute(
                "SELECT * FROM price_history ORDER BY observed_at DESC"
            ).fetchall()
        ]
        errors = [
            dict(row)
            for row in connection.execute(
                "SELECT * FROM scraping_errors ORDER BY created_at DESC"
            ).fetchall()
        ]
        geometry_metrics = [
            dict(row)
            for row in connection.execute(
                "SELECT * FROM geometry_metrics ORDER BY geometry_retrieved_at DESC"
            ).fetchall()
        ]
        evidence = [
            dict(row)
            for row in connection.execute(
                "SELECT * FROM evidence_records ORDER BY retrieved_at DESC"
            ).fetchall()
        ]
        provenance = [
            dict(row)
            for row in connection.execute(
                "SELECT * FROM data_provenance ORDER BY retrieved_at DESC"
            ).fetchall()
        ]
        source_run_status = [
            dict(row)
            for row in connection.execute(
                """
                SELECT * FROM source_run_status
                ORDER BY COALESCE(finished_at, started_at) DESC, source_id
                """
            ).fetchall()
        ]
        source_audit = [
            dict(row)
            for row in connection.execute(
                """
                SELECT source_record_id, job_id, source, source_type, external_id, title,
                       organization, organization_region, organization_commune,
                       organization_address, publication_date, closing_date, status,
                       region, commune, property_region, property_commune,
                       property_address, amount, currency, amount_raw, amount_type,
                       amount_source, amount_confidence, surface_raw, surface_m2,
                       surface_source, surface_confidence, source_url, detail_fetched,
                       detail_attempts, detail_http_status, fields_extracted_json,
                       fields_missing_json, extraction_reasons_json, data_completeness_score,
                       automatic_classification,
                       automatic_confidence, automatic_score, classification_reason,
                       positive_terms_json, negative_terms_json, rules_triggered_json,
                       manual_classification, manual_review_status, property_id, updated_at
                FROM source_records
                ORDER BY updated_at DESC, source_record_id
                """
            ).fetchall()
        ]
        latest_run_row = connection.execute(
            """
            SELECT * FROM scraping_runs
            ORDER BY COALESCE(finished_at, started_at, created_at) DESC
            LIMIT 1
            """
        ).fetchone()
        latest_run = dict(latest_run_row) if latest_run_row else None
        images = {
            row["internal_id"]: dict(row)
            for row in connection.execute("SELECT * FROM images").fetchall()
        }

    workbook = Workbook()
    summary = workbook.active
    summary.title = "Resumen"
    _write_summary(
        summary,
        properties,
        sources,
        source_audit,
        errors,
        latest_run,
        source_run_status,
    )
    terrenos = workbook.create_sheet("Terrenos")
    _write_properties(
        terrenos,
        properties,
        images,
        _group_by(sources, "internal_id"),
        _group_by(provenance, "property_id"),
        _group_by(source_audit, "property_id"),
    )
    _write_used_sources(workbook.create_sheet("Fuentes_Utilizadas"), source_run_status)
    _write_table(workbook.create_sheet("Fuentes"), sources)
    _write_table(workbook.create_sheet("Historial_Precios"), price_history)
    _write_table(workbook.create_sheet("Metricas_GIS"), geometry_metrics)
    _write_evidence(workbook.create_sheet("Fuentes_Evidencia"), evidence, provenance)
    _write_classifier_audit(workbook.create_sheet("Classifier_Audit"), source_audit)
    _write_source_records(workbook.create_sheet("SourceRecords"), source_audit)
    _write_table(workbook.create_sheet("Errores"), errors)

    workbook.save(output)
    load_workbook(output).close()
    return output


def _write_summary(
    sheet,
    properties: list[dict[str, Any]],
    sources: list[dict[str, Any]],
    source_audit: list[dict[str, Any]],
    errors: list[dict[str, Any]],
    latest_run: dict[str, Any] | None,
    source_run_status: list[dict[str, Any]],
) -> None:
    used_sources = [
        _source_label(row["source_id"])
        for row in source_run_status
        if row.get("status") in {"SUCCESS", "SUCCESS_WITH_WARNINGS"}
    ]
    omitted_sources = [
        f"{_source_label(row['source_id'])} - {row.get('status')}"
        for row in source_run_status
        if str(row.get("status", "")).startswith("SKIPPED")
    ]
    enrichment_used = [
        _source_label(row["source_id"])
        for row in source_run_status
        if row.get("source_type") == "EnrichmentSource"
        and row.get("status") in {"SUCCESS", "SUCCESS_WITH_WARNINGS"}
    ]
    rows = [
        ("Generado el", datetime.now(UTC).isoformat()),
        ("Version Landfinder", __version__),
        ("SearchJob ID", latest_run.get("run_id") if latest_run else "Sin job"),
        ("Fuentes solicitadas", _sources_from_run(latest_run)),
        (
            "Fuentes disponibles",
            ", ".join(sorted({_source_label(row["source_id"]) for row in source_run_status})),
        ),
        ("Fuentes utilizadas realmente", ", ".join(used_sources) or "Sin registros"),
        ("Fuentes omitidas", "; ".join(omitted_sources) or "Sin omisiones"),
        ("Servicios de enriquecimiento utilizados", ", ".join(enrichment_used) or "Sin registros"),
        ("Registros encontrados", latest_run.get("found_count") if latest_run else len(sources)),
        ("Registros fuente analizados", len(source_audit)),
        ("Registros exportados", len(properties)),
        ("Warnings", _count_json_messages(source_run_status, "warnings")),
        ("Errores", len(errors) + _count_json_messages(source_run_status, "errors")),
    ]
    for row_index, row in enumerate(rows, start=1):
        sheet.cell(row=row_index, column=1, value=row[0])
        sheet.cell(row=row_index, column=2, value=row[1])
    sheet.column_dimensions["A"].width = 24
    sheet.column_dimensions["B"].width = 32


def _write_properties(
    sheet,
    properties: list[dict[str, Any]],
    images: dict[str, dict[str, Any]],
    sources_by_property: dict[str, list[dict[str, Any]]],
    provenance_by_property: dict[str, list[dict[str, Any]]],
    source_audit_by_property: dict[str, list[dict[str, Any]]],
) -> None:
    headers = [
        "ID",
        "Imagen preview",
        "Titulo",
        "Fuente",
        "Fuente_Principal",
        "Fuentes_Encontradas",
        "APIs_Consultadas",
        "Servicios_Enriquecimiento",
        "URL",
        "Clasificacion",
        "Monto original",
        "Tipo de monto",
        "Region terreno",
        "Comuna terreno",
        "Direccion terreno",
        "Superficie m2",
        "Monto CLP",
        "Monto equivalente UF",
        "Fecha UF",
        "Tipo ubicacion",
        "Fuente ubicacion",
        "UF/m2",
        "Latitud",
        "Longitud",
        "Estado geometria",
        "Tipo geometria",
        "Fuente geometria",
        "Superficie GIS m2",
        "Perimetro GIS m",
        "Superficie oficial m2",
        "Diferencia area %",
        "first_seen",
        "last_seen",
        "Observaciones",
    ]
    sheet.append(headers)
    _style_header(sheet)
    sheet.freeze_panes = "A2"
    sheet.auto_filter.ref = f"A1:{get_column_letter(len(headers))}{max(1, len(properties) + 1)}"

    for row_index, item in enumerate(properties, start=2):
        internal_id = item["internal_id"]
        image = images.get(internal_id)
        property_sources = sources_by_property.get(internal_id, [])
        provenance = provenance_by_property.get(internal_id, [])
        audit = (source_audit_by_property.get(internal_id) or [{}])[0]
        values = [
            internal_id,
            "",
            item.get("publication_title"),
            item.get("source"),
            _source_label(item.get("source")),
            "; ".join(
                sorted({_source_label(row.get("source")) for row in property_sources})
            ),
            "; ".join(
                sorted(
                    {
                        _source_label(row.get("source_id"))
                        for row in provenance
                        if row.get("source_id")
                    }
                    | {_source_label(row.get("source")) for row in property_sources}
                )
            ),
            "; ".join(
                sorted(
                    {
                        _source_label(row.get("source_id"))
                        for row in provenance
                        if row.get("source_type") == "EnrichmentSource"
                    }
                )
            ),
            item.get("canonical_url"),
            audit.get("automatic_classification"),
            audit.get("amount_raw"),
            audit.get("amount_type"),
            item.get("region"),
            item.get("commune"),
            item.get("normalized_address") or item.get("original_address"),
            item.get("gross_area_m2"),
            item.get("asking_price_clp"),
            item.get("asking_price_uf"),
            item.get("uf_rate_date"),
            item.get("location_precision"),
            item.get("geocoding_source"),
            item.get("price_uf_m2_gross"),
            item.get("latitude"),
            item.get("longitude"),
            item.get("geometry_status"),
            item.get("geometry_type"),
            item.get("geometry_source"),
            item.get("geometry_area_m2"),
            item.get("geometry_perimeter_m"),
            item.get("official_area_m2"),
            item.get("area_difference_percent"),
            item.get("created_at"),
            item.get("last_seen"),
            item.get("anomaly_flags_json"),
        ]
        sheet.append(values)
        if item.get("canonical_url"):
            sheet.cell(row=row_index, column=9).hyperlink = item["canonical_url"]
            sheet.cell(row=row_index, column=9).style = "Hyperlink"
        preview_path = image.get("preview_path") if image else None
        if preview_path and Path(preview_path).exists():
            try:
                excel_image = ExcelImage(preview_path)
                excel_image.width = 144
                excel_image.height = 96
                sheet.add_image(excel_image, f"B{row_index}")
                sheet.row_dimensions[row_index].height = 78
            except Exception:
                sheet.cell(row=row_index, column=2, value=preview_path)

    widths = [
        38,
        18,
        40,
        20,
        22,
        30,
        36,
        36,
        34,
        24,
        18,
        18,
        22,
        20,
        36,
        15,
        16,
        14,
        12,
        12,
        12,
        18,
        16,
        22,
        18,
        16,
        18,
        16,
        24,
        24,
        28,
    ]
    for index, width in enumerate(widths, start=1):
        sheet.column_dimensions[get_column_letter(index)].width = width

    for column in (15, 16, 17, 21, 27, 28, 29, 30):
        for cell in sheet.iter_cols(min_col=column, max_col=column, min_row=2):
            for item in cell:
                item.number_format = "#,##0.00"


def _write_used_sources(sheet, rows: list[dict[str, Any]]) -> None:
    headers = [
        "Fuente",
        "Tipo",
        "Configurada",
        "Habilitada en busqueda",
        "Estado",
        "Uso",
        "Consultas",
        "Registros obtenidos",
        "Fecha ultima consulta",
    ]
    sheet.append(headers)
    _style_header(sheet)
    for row in rows:
        sheet.append(
            [
                _source_label(row.get("source_id")),
                row.get("source_type"),
                _yes_no(row.get("configured")),
                _yes_no(row.get("enabled")),
                row.get("status"),
                _source_usage(row.get("source_id")),
                "; ".join(_decode_list(row.get("queries_json"))),
                row.get("records_received"),
                row.get("finished_at") or row.get("started_at"),
            ]
        )
    if not rows:
        sheet.append(["Sin registros", "", "", "", "", "", "", "", ""])
    sheet.freeze_panes = "A2"
    sheet.auto_filter.ref = f"A1:I{max(2, len(rows) + 1)}"
    widths = [28, 20, 14, 24, 24, 24, 42, 20, 28]
    for index, width in enumerate(widths, start=1):
        sheet.column_dimensions[get_column_letter(index)].width = width


def _write_classifier_audit(sheet, rows: list[dict[str, Any]]) -> None:
    headers = [
        "SourceRecord_ID",
        "SearchJob_ID",
        "Fuente",
        "External_ID",
        "Titulo",
        "Organismo",
        "Clasificacion_Automatica",
        "Confianza",
        "Score",
        "Terminos_Positivos",
        "Terminos_Negativos",
        "Reglas",
        "Revision_Manual",
        "Estado_Revision",
        "Property_ID",
        "URL",
        "Razon",
        "Actualizado",
    ]
    sheet.append(headers)
    _style_header(sheet)
    for row_index, row in enumerate(rows, start=2):
        sheet.append(
            [
                row.get("source_record_id"),
                row.get("job_id"),
                _source_label(row.get("source")),
                row.get("external_id"),
                row.get("title"),
                row.get("organization"),
                row.get("automatic_classification"),
                row.get("automatic_confidence"),
                row.get("automatic_score"),
                "; ".join(_decode_list(row.get("positive_terms_json"))),
                "; ".join(_decode_list(row.get("negative_terms_json"))),
                "; ".join(_decode_list(row.get("rules_triggered_json"))),
                row.get("manual_classification"),
                row.get("manual_review_status"),
                row.get("property_id"),
                _sanitize_url(row.get("source_url")),
                row.get("classification_reason"),
                row.get("updated_at"),
            ]
        )
        if row.get("source_url"):
            sheet.cell(row=row_index, column=16).hyperlink = _sanitize_url(row["source_url"])
            sheet.cell(row=row_index, column=16).style = "Hyperlink"
    if not rows:
        sheet.append(
            ["Sin registros", "", "", "", "", "", "", "", "", "", "", "", "", "", "", "", "", ""]
        )
    sheet.freeze_panes = "A2"
    sheet.auto_filter.ref = f"A1:R{max(2, len(rows) + 1)}"
    widths = [30, 38, 22, 18, 48, 34, 26, 12, 10, 34, 34, 38, 22, 18, 38, 42, 56, 26]
    for index, width in enumerate(widths, start=1):
        sheet.column_dimensions[get_column_letter(index)].width = width


def _write_source_records(sheet, rows: list[dict[str, Any]]) -> None:
    headers = [
        "SourceRecord_ID",
        "SearchJob_ID",
        "Fuente",
        "External_ID",
        "Titulo",
        "Organismo",
        "Region organismo",
        "Comuna organismo",
        "Direccion organismo",
        "Region terreno",
        "Comuna terreno",
        "Direccion terreno",
        "Clasificacion",
        "Confianza",
        "Score",
        "Monto raw",
        "Monto CLP",
        "Tipo monto",
        "Superficie raw",
        "Superficie m2",
        "Detail fetched",
        "Detail attempts",
        "HTTP status",
        "Campos extraidos",
        "Campos faltantes",
        "Razones",
        "URL",
        "Property_ID",
        "Actualizado",
    ]
    sheet.append(headers)
    _style_header(sheet)
    for row_index, row in enumerate(rows, start=2):
        sheet.append(
            [
                row.get("source_record_id"),
                row.get("job_id"),
                _source_label(row.get("source")),
                row.get("external_id"),
                row.get("title"),
                row.get("organization"),
                row.get("organization_region"),
                row.get("organization_commune"),
                row.get("organization_address"),
                row.get("property_region") or row.get("region"),
                row.get("property_commune") or row.get("commune"),
                row.get("property_address"),
                row.get("automatic_classification"),
                row.get("automatic_confidence"),
                row.get("automatic_score"),
                row.get("amount_raw"),
                row.get("amount"),
                row.get("amount_type"),
                row.get("surface_raw"),
                row.get("surface_m2"),
                _yes_no(row.get("detail_fetched")),
                row.get("detail_attempts"),
                row.get("detail_http_status"),
                "; ".join(_decode_list(row.get("fields_extracted_json"))),
                "; ".join(_decode_list(row.get("fields_missing_json"))),
                _json_value(row.get("extraction_reasons_json")),
                _sanitize_url(row.get("source_url")),
                row.get("property_id"),
                row.get("updated_at"),
            ]
        )
        if row.get("source_url"):
            sheet.cell(row=row_index, column=27).hyperlink = _sanitize_url(row["source_url"])
            sheet.cell(row=row_index, column=27).style = "Hyperlink"
    if not rows:
        sheet.append(["Sin registros"] + [""] * (len(headers) - 1))
    sheet.freeze_panes = "A2"
    sheet.auto_filter.ref = f"A1:{get_column_letter(len(headers))}{max(2, len(rows) + 1)}"
    widths = [
        30,
        38,
        22,
        18,
        48,
        34,
        24,
        22,
        34,
        22,
        22,
        34,
        24,
        12,
        10,
        22,
        16,
        20,
        18,
        16,
        14,
        14,
        12,
        34,
        34,
        48,
        42,
        38,
        26,
    ]
    for index, width in enumerate(widths, start=1):
        sheet.column_dimensions[get_column_letter(index)].width = width
    for column in (14, 17, 20):
        for cell in sheet.iter_cols(min_col=column, max_col=column, min_row=2):
            for item in cell:
                item.number_format = "#,##0.00"


def _write_evidence(
    sheet,
    evidence: list[dict[str, Any]],
    provenance: list[dict[str, Any]],
) -> None:
    headers = [
        "Property_ID",
        "Campo",
        "Valor",
        "Fuente",
        "Tipo_Fuente",
        "Endpoint_Sanitizado",
        "Fecha_Consulta",
        "Metodo",
        "Confianza",
    ]
    sheet.append(headers)
    _style_header(sheet)
    for row in provenance:
        sheet.append(
            [
                row.get("property_id"),
                row.get("field_name"),
                _json_value(row.get("value")),
                _source_label(row.get("source_id")),
                row.get("source_type"),
                _sanitize_url(row.get("source_url")),
                row.get("retrieved_at"),
                row.get("method"),
                row.get("confidence"),
            ]
        )
    for row in evidence:
        source = row.get("source_name") or row.get("source")
        sheet.append(
            [
                row.get("property_id") or row.get("internal_id"),
                row.get("field"),
                _json_value(row.get("value")),
                source,
                _source_type_for_evidence(source),
                _sanitize_url(row.get("source_url") or row.get("url")),
                row.get("retrieved_at"),
                row.get("method"),
                row.get("confidence"),
            ]
        )
    total = len(provenance) + len(evidence)
    if total == 0:
        sheet.append(["Sin registros", "", "", "", "", "", "", "", ""])
    sheet.freeze_panes = "A2"
    sheet.auto_filter.ref = f"A1:I{max(2, total + 1)}"
    widths = [38, 28, 42, 28, 20, 48, 28, 18, 14]
    for index, width in enumerate(widths, start=1):
        sheet.column_dimensions[get_column_letter(index)].width = width


def _write_table(sheet, rows: list[dict[str, Any]]) -> None:
    if not rows:
        sheet.append(["Sin registros"])
        return
    headers = list(rows[0].keys())
    sheet.append(headers)
    _style_header(sheet)
    for row in rows:
        sheet.append([row.get(header) for header in headers])
    sheet.freeze_panes = "A2"
    sheet.auto_filter.ref = f"A1:{get_column_letter(len(headers))}{len(rows) + 1}"
    for index, header in enumerate(headers, start=1):
        width = min(max(len(header) + 4, 14), 42)
        sheet.column_dimensions[get_column_letter(index)].width = width


def _group_by(rows: list[dict[str, Any]], key: str) -> dict[str, list[dict[str, Any]]]:
    result: dict[str, list[dict[str, Any]]] = {}
    for row in rows:
        value = row.get(key)
        if value is None:
            continue
        result.setdefault(str(value), []).append(row)
    return result


def _source_label(source_id: Any) -> str:
    return {
        "mercado_publico": "Mercado Publico",
        "mercadolibre_real_estate": "MercadoLibre",
        "cmf_uf": "CMF",
        "minvu": "IDE MINVU",
    }.get(str(source_id), "" if source_id is None else str(source_id))


def _source_usage(source_id: Any) -> str:
    return {
        "mercado_publico": "Descubrimiento",
        "mercadolibre_real_estate": "Publicaciones",
        "cmf_uf": "Conversion UF",
        "minvu": "Territorial",
    }.get(str(source_id), "")


def _yes_no(value: Any) -> str:
    if value in (1, True, "1", "true", "True"):
        return "Si"
    if value in (0, False, "0", "false", "False"):
        return "No"
    return "-"


def _decode_list(value: Any) -> list[str]:
    if value in (None, ""):
        return []
    if isinstance(value, list):
        return [str(item) for item in value]
    try:
        decoded = json.loads(str(value))
    except json.JSONDecodeError:
        return [str(value)]
    if isinstance(decoded, list):
        return [str(item) for item in decoded]
    return [str(decoded)]


def _sources_from_run(run: dict[str, Any] | None) -> str:
    if not run:
        return "Sin job"
    return ", ".join(_source_label(item) for item in _decode_list(run.get("sources_json")))


def _count_json_messages(rows: list[dict[str, Any]], key: str) -> int:
    return sum(len(_decode_list(row.get(key))) for row in rows)


def _json_value(value: Any) -> str:
    if value is None:
        return ""
    if isinstance(value, str):
        try:
            decoded = json.loads(value)
        except json.JSONDecodeError:
            return value
        return json.dumps(decoded, ensure_ascii=True)
    return json.dumps(value, ensure_ascii=True, default=str)


def _sanitize_url(value: Any) -> str:
    if value is None:
        return ""
    return re.sub(
        r"(?i)(ticket|apikey|access_token|refresh_token)=([^&\s]+)",
        r"\1=REDACTED",
        str(value),
    )


def _source_type_for_evidence(source: Any) -> str:
    text = str(source or "").lower()
    if "minvu" in text or "cmf" in text:
        return "EnrichmentSource"
    return "SearchSource"


def _style_header(sheet) -> None:
    fill = PatternFill("solid", fgColor="1F2937")
    for cell in sheet[1]:
        cell.font = Font(color="FFFFFF", bold=True)
        cell.fill = fill
